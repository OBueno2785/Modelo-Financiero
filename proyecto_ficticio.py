"""Prueba de generalidad: una PTAR inventada, sin una sola cifra de Cajamarca.

La razon de ser del generador es que la opcion "Proyecto nuevo" produzca los
mismos flujos que los modelos aprobados. Validar contra Cajamarca y San Martin
no lo demuestra: los motores nacieron de esos dos libros y pueden traer
supuestos suyos pegados. Esta plantilla es un tercer proyecto, con otras fechas,
otro plazo, otro apalancamiento y otros costos, y sirve de control: si el
generador solo sabe hacer Cajamarca, aca se cae.

Escribe `Plantilla_PTAR_Ficticia.xlsx`. No toca ningun archivo de Oscar.
"""
from datetime import date
from openpyxl import load_workbook

RUTA_BASE = "/mnt/project-files/generador/Plantilla_Proyecto_Nuevo.xlsx"
RUTA = "/mnt/project-files/generador/Plantilla_PTAR_Ficticia.xlsx"

IPC = 0.02


def _fin_de_mes(a, m):
    return (date(a + (m == 12), m % 12 + 1, 1) - __import__("datetime").timedelta(1)).isoformat()


MALLA = [_fin_de_mes(a, m) for a in range(2027, 2053) for m in range(1, 13)]

CIERRE      = "2027-01-31"
CIERRE_FIN  = "2028-03-31"
INI_OBRA    = "2028-04-30"
FIN_OBRA    = "2031-03-31"
INI_OP      = "2031-03-31"
FIN_CONC    = "2051-03-31"
INI_DEUDA   = "2031-04-30"
FIN_DEUDA   = "2043-03-31"

# Precios constantes de diciembre de 2025, igual que la plantilla base, para que
# los dos indices sean meses enteros de IPC y se puedan verificar a mano.
MESES_A_OM    = 63   # dic-2025 -> mar-2031
MESES_A_CIERRE = 13  # dic-2025 -> ene-2027

PROYECTO = {
    "Nombre del proyecto": "PTAR Ficticia",
    "Fecha de cierre del contrato": CIERRE,
    "Fecha de cierre financiero": CIERRE_FIN,
    "Inicio de obras": INI_OBRA,
    "Fin de obras y puesta en marcha": FIN_OBRA,
    "Inicio de operacion": INI_OP,
    "Fin de la concesion": FIN_CONC,
    "Regimen de la APP": "cofinanciada",
    "Pago del PPDI": "al final de obra",
    "Numero de cuotas del PPDI": 48,
    "Factor de conversion a kg DBO5": 0.38,
    "Tamano de deuda": 0.70,
    "Spread sobre el soberano": 0.04,
    "Inicio del pago de deuda": INI_DEUDA,
    "Fin del pago de deuda": FIN_DEUDA,
    "Costo de estructuracion": 0.015,
    "Comision por no uso": 0.0075,
    "Cobertura minima RCSD": 1.25,
    "Cuenta de reserva CRSD": 0.5,
    "IPM, reajuste de ingresos": 0.0267,
    "Gatillo del reajuste": 0.03,
    "IPC, reajuste de costos": IPC,
    "Indice de costos al inicio de O&M": (1 + IPC) ** (MESES_A_OM / 12),
    "Indice de CAPEX al cierre del contrato": (1 + IPC) ** (MESES_A_CIERRE / 12),
    "IGV": 0.18,
    "Impuesto a la Renta": 0.295,
    "Participacion de trabajadores": 0.05,
    "Margen del costo fijo": 0.15,
    "Supervision del regulador": 0.01,
}


def capex():
    """S/ miles a precios constantes. Estudios y terrenos antes de la obra, y
    una obra de 35 meses con perfil de campana, distinta del perfil cargado al
    frente de Cajamarca."""
    v = {f: 0.0 for f in MALLA}
    for f in MALLA[1:MALLA.index(INI_OBRA) + 1]:      # feb-2027 .. abr-2028
        v[f] = 420.0
    i0, i1 = MALLA.index(INI_OBRA), MALLA.index(FIN_OBRA)
    vent = MALLA[i0 + 1:i1 + 1]                        # lo que financia la linea
    n = len(vent)
    peso = [1 - abs((k - (n - 1) / 2) / ((n - 1) / 2)) * 0.6 for k in range(n)]
    tot = sum(peso)
    for k, f in enumerate(vent):
        v[f] = 180_000.0 * peso[k] / tot
    return [v[f] for f in MALLA]


ANIO_CON = list(range(5, 26))                    # 21 anos de operacion
ANIO_CAL = [2031 + i for i in range(len(ANIO_CON))]
NA = len(ANIO_CON)

DEMANDA = [8_400_000 * 1.012 ** i for i in range(NA)]


def opex():
    esc = [d / DEMANDA[0] for d in DEMANDA]
    rep = [0.0] * NA
    for a in (9, 14, 19):                            # reposiciones cada 5 anos
        rep[a] = 3_600_000.0
    return {
        "personal":            [2_050_000.0] * NA,
        "mant_obra_civil":     [  480_000.0] * NA,
        "mant_colectores":     [  310_000.0] * NA,
        "energia_electrica":   [  920_000.0] * NA,
        "productos_materiales":[  660_000.0] * NA,
        "gestion_monitoreo":   [  340_000.0] * NA,
        "reposiciones":        rep,
        "energia_ptar":        [1_750_000.0 * e for e in esc],
        "reactivos":           [  610_000.0 * e for e in esc],
        "residuos":            [  395_000.0 * e for e in esc],
        "gastos_spv":          [  520_000.0] * NA,
        "demanda_m3":          DEMANDA,
    }


def otros():
    """S/ miles. Bloque de construccion y bloque de operacion."""
    cero = {f: 0.0 for f in MALLA}
    con = {c: dict(cero) for c in
           ("seguros", "fianzas", "fideicomiso", "promocion", "garantias", "reembolso")}
    con["fideicomiso"][CIERRE_FIN] = 260.0
    con["reembolso"][CIERRE] = 6_100.0
    con["garantias"][CIERRE_FIN] = 480.0
    op = {c: dict(cero) for c in ("seguros", "fianzas", "fideicomiso", "promocion")}
    i0, i1 = MALLA.index(INI_OP), MALLA.index(FIN_CONC)
    for f in MALLA[i0 + 1:i1 + 1]:
        op["seguros"][f] = 18.0
        op["fianzas"][f] = 6.5
        op["fideicomiso"][f] = 2.0
    return con, op


# ------------------------------------------------- esquema por hitos
ANIO0_H, NM_H = 2027, 312
NY_H = NM_H // 12
PCT_HITO = [0.45, 0.20, 0.20, 0.15]

HITOS = {
    "Primer ano de la malla": ANIO0_H,
    "Meses de la malla": NM_H,
    "Inicio del periodo D + C + PM": "2027-02",
    "Fin del periodo D + C + PM": "2031-03",
    "Primer pago del PPDI": "2031-06",
    "Fin de la operacion": "2051-03",
    "Mes del cierre financiero": 15,
    "Apalancamiento": 0.70,
    "Comision de estructuracion": 0.015,
    "Comision por no uso": 0.0075,
    "Aporte 1, porcentaje": 0.25, "Aporte 1, mes del contrato": 1,
    "Aporte 2, porcentaje": 0.25, "Aporte 2, mes del contrato": 10,
    "Aporte 3, porcentaje": 0.50, "Aporte 3, mes del contrato": 16,
    "Capital de trabajo inicial": 1_200.0,
    "Periodo medio de pago a proveedores": 1.0,
    "Sistema 1, kg DBO al ano": 0.9 * DEMANDA[0] * PROYECTO["Factor de conversion a kg DBO5"],
    "Sistema 2, kg DBO al ano": 0.1 * DEMANDA[0] * PROYECTO["Factor de conversion a kg DBO5"],
}


def opex_hitos():
    """Por ano calendario de la malla (2027..2052), en S/ miles constantes.

    Las lineas de O&M van en cero fuera de los anos de operacion, igual que el
    libro de San Martin. El motor las prorratea con la bandera de operacion, asi
    que dejarlas llenas no cambiaria ningun resultado, pero rompe la regla de
    Oscar de que el OPEX por hito sume el OPEX del ano: la hoja `OPEX` no cubre
    esos anos. Lo que escriba esta funcion es solo el punto de partida, porque
    `opex_consistente.derivar` reescribe las patas por hito desde `OPEX`.
    """
    conv = PROYECTO["Factor de conversion a kg DBO5"]
    esc = lambda i: 1.012 ** max(0, ANIO0_H + i - 2031)
    op = lambda i: 1.0 if 2031 <= ANIO0_H + i <= 2051 else 0.0
    fijo = 4_760.0                      # las seis lineas fijas, sin reposiciones
    var = 2_755.0                       # energia de PTAR, reactivos y residuos
    L = {}
    for k, pct in zip(("fijo_h1", "fijo_h2", "fijo_h3", "fijo_h4"), PCT_HITO):
        L[k] = [fijo * pct * op(i) for i in range(NY_H)]
        L["conigv_h" + k[-1]] = [fijo * pct * 1.15 * op(i) for i in range(NY_H)]
    L["variable"] = [var * esc(i) * op(i) for i in range(NY_H)]
    L["conigv_var_tar"] = [0.9 * var * esc(i) * 1.15 * op(i) for i in range(NY_H)]
    L["conigv_var_sjs"] = [0.1 * var * esc(i) * 1.15 * op(i) for i in range(NY_H)]
    L["gg_u"] = [800.0 * op(i) for i in range(NY_H)]
    L["otros_op"] = [450.0 * op(i) for i in range(NY_H)]
    L["igv_op"] = [300.0 * op(i) for i in range(NY_H)]
    rep = [0.0] * NY_H
    for a in (2040, 2045, 2050):
        rep[a - ANIO0_H] = 3_600.0
    L["repex"] = rep
    L["pago_obras"] = [0.0] * NY_H
    L["cierre_laguna"] = [0.0] * NY_H
    L["carga_1"] = [0.9 * DEMANDA[0] * conv * esc(i) for i in range(NY_H)]
    L["carga_2"] = [0.1 * DEMANDA[0] * conv * esc(i) for i in range(NY_H)]
    return L


def _hitos(wb):
    from plantilla import LINEAS_OPEX_HITOS
    ws = wb["Hitos"]
    vistos = set()
    fila_pct = []
    for r in range(6, ws.max_row + 1):
        et = ws.cell(r, 2).value
        if et in HITOS:
            ws.cell(r, 4).value = HITOS[et]
            vistos.add(et)
        elif isinstance(et, str) and et.startswith("Hito "):
            fila_pct.append(r)
    faltan = set(HITOS) - vistos
    if faltan:
        raise SystemExit(f"la hoja Hitos no tiene: {sorted(faltan)}")
    for r, pct in zip(fila_pct, PCT_HITO):
        ws.cell(r, 4).value = pct

    ws = wb["OPEX hitos"]
    for r in range(5, ws.max_row + 1):
        for c in range(4, ws.max_column + 1):
            ws.cell(r, c).value = None
    for j in range(NY_H):
        ws.cell(5, 4 + j).value = ANIO0_H + j
    L = opex_hitos()
    clave_de = {e: c for e, c in LINEAS_OPEX_HITOS if c}
    for r in range(7, ws.max_row + 1):
        c = clave_de.get(ws.cell(r, 2).value)
        if c:
            for j in range(NY_H):
                ws.cell(r, 4 + j).value = L[c][j]


def escribir(ruta=RUTA, regimen="cofinanciada", esquema="al final de obra"):
    from plantilla import LINEAS_OPEX, OTROS, OTROS_OP
    wb = load_workbook(RUTA_BASE)

    ws = wb["Proyecto"]
    P = dict(PROYECTO, **{"Regimen de la APP": regimen,
                          "Pago del PPDI": esquema})
    vistos = set()
    for r in range(6, ws.max_row + 1):
        et = ws.cell(r, 2).value
        if et in P:
            ws.cell(r, 4).value = P[et]
            vistos.add(et)
    faltan = set(P) - vistos
    if faltan:
        raise SystemExit(f"la plantilla no tiene estas filas: {sorted(faltan)}")

    # El `Cronograma` viene con las etapas de Cajamarca y sus fechas MANDAN
    # sobre las de `Proyecto`, asi que hay que reescribirlo con las de este
    # proyecto o la plantilla queda corriendo el calendario del libro aprobado.
    from plantilla import volcar_cronograma
    volcar_cronograma(wb, P, MALLA)

    ws = wb["CAPEX"]
    for r in range(6, ws.max_row + 2):               # borra la malla de Cajamarca
        ws.cell(r, 2).value = ws.cell(r, 4).value = None
    cx = capex()
    for i, f in enumerate(MALLA):
        ws.cell(6 + i, 2).value = f
        ws.cell(6 + i, 4).value = cx[i]

    ws = wb["OPEX"]
    for r in range(5, ws.max_row + 1):
        for c in range(4, ws.max_column + 1):
            ws.cell(r, c).value = None
    for j in range(NA):
        ws.cell(5, 4 + j).value = ANIO_CON[j]
        ws.cell(6, 4 + j).value = ANIO_CAL[j]
    ox = opex()
    clave_de = {e: c for e, c in LINEAS_OPEX if c}
    for r in range(8, ws.max_row + 1):
        c = clave_de.get(ws.cell(r, 2).value)
        if c:
            for j in range(NA):
                ws.cell(r, 4 + j).value = ox[c][j]

    ws = wb["Otros"]
    for r in range(6, ws.max_row + 2):
        for c in range(2, ws.max_column + 1):
            ws.cell(r, c).value = None
    oc, oo = otros()
    for i, f in enumerate(MALLA):
        ws.cell(6 + i, 2).value = f
        for j, (_, k) in enumerate(OTROS):
            ws.cell(6 + i, 4 + j).value = oc[k][f]
        for j, (_, k) in enumerate(OTROS_OP):
            ws.cell(6 + i, 4 + len(OTROS) + j).value = oo[k][f]

    _hitos(wb)
    wb.save(ruta)
    # La regla de Oscar: `OPEX hitos` se deriva de `OPEX` mas el reparto por
    # hito, para que el usuario escriba el OPEX una sola vez.
    import opex_consistente as OC
    OC.derivar(ruta)
    return ruta


def control(ipm=0.0267):
    """Las cuatro combinaciones sobre el proyecto inventado, desde la plantilla.

    Es el gemelo de `matriz.py`: aquella corre los dos proyectos aprobados, esta
    corre uno que no existe. Las dos tienen que cerrar en VAN cero.
    """
    import json
    import generar as G
    import generar_sm as GS
    A = json.load(open("/mnt/project-files/generador/cajamarca_actualizado.json"))["hoy"]
    out = []

    g = G.correr(escribir(RUTA, "cofinanciada"),
                 kd=A["kd"], ke=A["ke"], wacc=A["wacc"])
    out.append(("cofinanciada", "al final de obra", g["r"]["van"],
                f"PPDI {g['ppdi']:,.2f} / tarifa {g['tarifa']:.7f}"))

    g = G.correr_auto(escribir(_r("auto"), "autofinanciada"),
                      kd=A["kd"], ke=A["ke"], wacc=A["wacc"])
    out.append(("autofinanciada", "al final de obra", g["r"]["van"],
                f"tarifa {g['tarifa']:.7f} / intangible "
                f"{g['r']['intangible']['base']:,.2f}"))

    kw = dict(kd=A["kd"], wacc=A["wacc"], tir=round(A["ke"], 4), ipm=ipm,
              tasa_reparto=A["wacc"])
    g = GS.correr_plantilla(escribir(_r("hitos_cofi"), "cofinanciada", "por hitos"), **kw)
    out.append(("cofinanciada", "por hitos", g["r"]["van"],
                f"PPDI {g['ppdi']:,.2f} / referencia {g['ref']:,.4f}"))

    g = GS.correr_plantilla(escribir(_r("hitos_auto"), "autofinanciada", "por hitos"), **kw)
    out.append(("autofinanciada", "por hitos", g["r"]["van"],
                f"tarifa {g['r']['precio']:.7f} / intangible "
                f"{g['r']['intangible']['base']:,.2f}"))
    return out


def _r(sufijo):
    return RUTA.replace(".xlsx", f"_{sufijo}.xlsx")


if __name__ == "__main__":
    print("escrita:", escribir())
    import leer_plantilla as LP
    d = LP.leer(RUTA)
    s = LP.series(d)
    c = LP.construccion(d)
    print(f"meses de la malla     : {len(d['fechas'])}   anos de OPEX: {d['anios']}")
    print(f"CAPEX constante       : {sum(d['capex']):>16,.2f} S/ miles")
    print(f"  de obra (financiado): {sum(c['inversion']):>16,.2f} corrientes")
    print(f"CAPEX corriente total : {sum(c['costo_obra']):>16,.2f}")
    print(f"PPDMO fijo base       : {s['fijo_base']:>16,.4f} S/ miles al mes")
    print(f"costo fijo total      : {sum(s['costo_fijo']):>16,.2f}")
    print(f"produccion total (kg) : {sum(s['kg']):>16,.0f}")
    print()
    print(f"{'regimen':<16}{'esquema':<18}{'VAN del cierre':>18}   resultado")
    for reg, esq, van, txt in control():
        print(f"{reg:<16}{esq:<18}{van:>18.8f}   {txt}")


def libros(ipm=0.0267, fx=0.0):
    """Escribe los cuatro libros de salida del proyecto inventado.

    El gemelo de lo que `matriz.correr_todo` hace con los dos aprobados: el
    control no sirve de nada si solo cierra en memoria y nadie abre el Excel,
    que es donde aparecieron la referencia circular de `Deuda` y la trampa de
    unidades del `Resumen`.
    """
    import json
    import escribir_excel as EX
    import escribir_excel_sm as EXS
    import generar as G
    import generar_sm as GS
    A = json.load(open("/mnt/project-files/generador/cajamarca_actualizado.json"))["hoy"]
    kw = dict(kd=A["kd"], ke=A["ke"], wacc=A["wacc"])
    sal = "/mnt/project-files/generador/"
    out = []

    g = G.correr(escribir(RUTA, "cofinanciada"), **kw)
    EX.escribir(G.payload(g), sal + "PTAR_Ficticia_generado.xlsx")
    out.append("PTAR_Ficticia_generado.xlsx")

    g = G.correr_auto(escribir(_r("auto"), "autofinanciada"), **kw)
    EX.escribir_auto(G.payload_auto(g), sal + "PTAR_Ficticia_autofinanciada.xlsx")
    out.append("PTAR_Ficticia_autofinanciada.xlsx")

    kwh = dict(kd=A["kd"], wacc=A["wacc"], tir=round(A["ke"], 4), ipm=ipm,
               tasa_reparto=A["wacc"])
    g = GS.correr_plantilla(escribir(_r("hitos_cofi"), "cofinanciada", "por hitos"), **kwh)
    EXS.escribir(GS.payload(g, A["ke"], A["kd"], A["wacc"], ipm, fx),
                 sal + "PTAR_Ficticia_hitos.xlsx")
    out.append("PTAR_Ficticia_hitos.xlsx")

    g = GS.correr_plantilla(escribir(_r("hitos_auto"), "autofinanciada", "por hitos"), **kwh)
    EXS.escribir_auto(GS.payload_auto(g, A["ke"], A["kd"], A["wacc"], ipm, fx),
                      sal + "PTAR_Ficticia_hitos_autofinanciada.xlsx")
    out.append("PTAR_Ficticia_hitos_autofinanciada.xlsx")
    return out
