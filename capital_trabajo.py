"""Facilidad de capital de trabajo de OPERACION (`M.Sombra!111:129` y `155:186`).

No es el tramo complementario de IGV, que en Cajamarca esta cableado y nunca
gira: lo que si gira es esta linea, y es de operacion, no de obra. Entra BRUTA
en el bloque de financiamiento del flujo de caja financiero (`M.Sombra!31` y `32`),
nunca neteada contra el servicio.

La regla de dimensionamiento es del libro y se reproduce sin leer ninguna cifra:

    necesidad del trimestre = (costos fijos + costos variables) x hito de O&M
    `M.Sombra!123 = SUM(D121:D122) * D114`

El interes de los trimestres intermedios se CAPITALIZA, igual que el de obra, y
el repago es BULLET: `Inputs!D254 = Trimestres de Repago = 1` significa que se
paga en un solo trimestre, no que venza al siguiente.

Cuantos giros hay NO es una regla general: el libro tiene tres bloques de tramo
cableados y solo dos con formula de desembolso, asi que financia los dos primeros
trimestres de operacion y despues asume que la operacion se autofinancia. La
necesidad sigue calculandose y creciendo todos los trimestres sin que nadie la
gire. Por eso el numero de giros es un INPUT con el 2 del libro por defecto, y no
algo derivado: derivarlo seria inventar una regla que el libro no tiene.
"""

TET_LIBRO = 0.016556249      # M.Sombra!D150, soberano a vida media + spread
GIROS_LIBRO = 2              # bloques con formula de desembolso
DESFASE_REPAGO = 3           # gira en q, capitaliza q+1 y q+2, repaga en q+3


def corrida(costo_fijo, costo_var, flag_om, nq, tet=TET_LIBRO,
            giros=GIROS_LIBRO, desfase=DESFASE_REPAGO, cobertura=1.0):
    """Devuelve las series trimestrales brutas de la facilidad.

    `costo_fijo` y `costo_var` son POSITIVOS (costo), `flag_om` el hito de O&M.
    `cobertura` es `Inputs!D250`, la fraccion de la necesidad que se financia.
    """
    desemb = [0.0] * nq
    servicio = [0.0] * nq
    interes = [0.0] * nq
    girados = 0
    for q in range(nq):
        if girados >= giros or not flag_om[q]:
            continue
        nec = (costo_fijo[q] + costo_var[q]) * cobertura
        if nec <= 0:
            continue
        girados += 1
        desemb[q] = nec
        saldo = nec
        for k in range(1, desfase):          # capitaliza los intermedios
            i = saldo * tet
            interes[q + k] += i
            saldo += i
        i_final = saldo * tet
        interes[q + desfase] += i_final
        servicio[q + desfase] += saldo + i_final
    return dict(desembolso=desemb, servicio=servicio, interes=interes)


if __name__ == "__main__":
    # Control contra Cajamarca: la necesidad se reconstruye de los costos del
    # libro, no se lee, y tiene que dar los dos giros al centimo.
    import openpyxl
    v = openpyxl.load_workbook("/mnt/project-files/generador/"
                               "MODEF_PTAR_Cajamarca.xlsm", data_only=True)["M.Sombra"]
    col = lambda r, n: [abs(v.cell(row=r, column=c).value or 0) for c in range(4, 4 + n)]
    n = 40
    cf, cv = col(121, n), col(122, n)
    fl = [1 if (v.cell(row=114, column=c).value or 0) else 0 for c in range(4, 4 + n)]
    r = corrida(cf, cv, fl, n)
    g = [(i, x) for i, x in enumerate(r["desembolso"]) if x]
    print(f"{'':<22}{'motor':>12}{'libro':>12}{'desvio':>10}")
    for (i, x), esp in zip(g, (6314.54, 6392.03)):
        print(f"  giro trim. {i+1:<10}{x:>12,.2f}{esp:>12,.2f}{x-esp:>10.2f}")
    tot_i, tot_s = sum(r["interes"]), sum(r["servicio"])
    print(f"  {'desembolsado':<20}{sum(r['desembolso']):>12,.2f}{12706.57:>12,.2f}"
          f"{sum(r['desembolso'])-12706.57:>10.2f}")
    print(f"  {'interes':<20}{tot_i:>12,.2f}{641.63:>12,.2f}{tot_i-641.63:>10.2f}")
    print(f"  {'servicio':<20}{tot_s:>12,.2f}{13348.20:>12,.2f}{tot_s-13348.20:>10.2f}")
