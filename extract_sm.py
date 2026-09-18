import openpyxl, json, warnings
warnings.filterwarnings("ignore")
from openpyxl.utils import column_index_from_string as ci
f="64199c72-Modelo_Hitos_Funcionales.xlsm"
wb = openpyxl.load_workbook(f, data_only=True)
c0, c1 = ci("F"), ci("MC")
def row(sh, r):
    ws = wb[sh]
    return [float(ws.cell(r,c).value) if isinstance(ws.cell(r,c).value,(int,float)) else 0.0
            for c in range(c0, c1+1)]
S = {
 "nec_total":      row("Nec_Fin",21),
 "nec_sin_crsd":   row("Nec_Fin",23),
 "nec_capex":      row("Nec_Fin",14),
 "nec_otros":      row("Nec_Fin",15),
 "nec_igv":        row("Nec_Fin",16),
 "nec_intereses":  row("Nec_Fin",17),
 "nec_comisiones": row("Nec_Fin",18),
 "nec_crsd":       row("Nec_Fin",19),
 "nec_ktrabajo":   row("Nec_Fin",20),
 "ppdi_trim":      row("Ingresos",38),
 "flag_pago":      row("Ingresos",34),
 "flag_dcpm":      row("Nec_Fin",6),
 "nro_mes":        row("Nec_Fin",7),
 "ppdmo_trim_tot": [0.0]*(c1-c0+1),
}
for r in (131,132,133,134,135):
    v = row("Ingresos", r)
    S["ppdmo_trim_tot"] = [a+b for a,b in zip(S["ppdmo_trim_tot"], v)]
K = {
 "col_F": c0, "col_P": ci("P"), "n": c1-c0+1,
 "wacc_anual":  wb["Ke"]["C25"].value,
 "ke_anual":    wb["Ke"]["C23"].value,
 "ke_usd":      wb["Ke"]["C14"].value,
 "rf":          wb["Ke"]["C3"].value,
 "erp":         wb["Ke"]["C6"].value,
 "rp":          wb["Ke"]["C4"].value,
 "beta_u":      wb["Ke"]["C5"].value,
 "w_deuda":     wb["Ke"]["C8"].value,
 "tasa_imp":    wb["Ke"]["C10"].value,
 "zc_sol":      wb["Ke"]["C21"].value,
 "zc_usd":      wb["Ke"]["C20"].value,
 "kd":          wb["Ke"]["C22"].value,
 "soberano":    wb["Tasa de Interés"]["D6"].value,
 "spread":      wb["Tasa de Interés"]["D10"].value,
 "van_nec":     wb["Nec_Fin"]["D24"].value,
 "van_ppdi":    wb["Ingresos"]["D47"].value,
 "ppdi_anual":  wb["Control"]["C16"].value,
 "ppdi_trim":   wb["Ingresos"]["D22"].value,
 "n_cuotas":    wb["Control"]["C15"].value,
 "ref_ppdmo":   wb["Control"]["C49"].value,
 "ppdmo_fijo_anual": wb["Control"]["C21"].value,
 "precio_unit": wb["Ingresos"]["D109"].value,
 "apalancamiento": wb["Control"]["F22"].value,
 "upfront": wb["Control"]["J12"].value,
 "commitment": wb["Control"]["J13"].value,
 "mes_cf": wb["Control"]["J5"].value,
 "ap_pct": [wb["Control"][f"F{r}"].value for r in (29,30,31)],
 "ap_mes": [wb["Nec_Fin"][f"D{r}"].value for r in (33,35,37)],
 "nec_total_libro": wb["Nec_Fin"]["D21"].value,
 "nec_int_libro": wb["Nec_Fin"]["D17"].value,
 "nec_com_libro": wb["Nec_Fin"]["D18"].value,
 "deuda_libro": wb["Nec_Fin"]["D48"].value,
 "capital_libro": wb["Nec_Fin"]["D39"].value,
 "pct_capex":   [wb["Capex"][f"D{r}"].value for r in (160,161,162,163)],
 "ppdi_hito":   [wb["Ingresos"][f"D{r}"].value for r in (25,26,27,28)],
 "ppdmo_hito":  [wb["Ingresos"][f"D{r}"].value for r in (94,95,96,97)],
}
json.dump({"series":S,"escalares":K}, open("sanmartin_primitivos.json","w"))
for k,v in K.items():
    print(f"{k:18} {v}")
